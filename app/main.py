"""Xodimlar kelish-ketish nazorati — asosiy ilova."""
import base64
import io
import math
import os
import secrets
import socket
from contextlib import asynccontextmanager
from datetime import datetime, date, timedelta, time as dtime

import qrcode
from fastapi import FastAPI, Form, Request
from fastapi.responses import (
    HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse,
)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from itsdangerous import BadSignature, URLSafeSerializer

from .config import (
    BASE_DIR, TZ, QR_ROTATE_SECONDS, QR_SLOTS_AHEAD, STATUS_LABELS,
    ABSENCE_TYPES, WEEKDAYS_UZ, MONTHS_UZ,
)
from . import db as D
from .security import make_token, verify_token, seconds_left, current_window


IS_CLOUD = bool(os.environ.get("VERCEL") or D.IS_PG)


@asynccontextmanager
async def lifespan(_app: FastAPI):
    """Ishga tushish.

    Bu yerda hech qachon xato ko'tarilmasligi kerak: serverless muhitda
    lifespan qulasa, butun funksiya ishlamay qoladi va foydalanuvchi
    sababini bilmaydigan 500 xatosini ko'radi. Baza esa birinchi so'rovda
    o'zi tayyorlanadi (`D.db()` buni o'zi qiladi).
    """
    if IS_CLOUD:
        # Bulutda manzil so'rov sarlavhalaridan olinadi, lokal IP ma'nosiz.
        # Bazaga ulanish ham shu yerda emas, birinchi so'rovda amalga oshadi.
        yield
        return

    try:
        D.ensure_db()
        ip = lan_ip()
        with D.db() as conn:
            if not D.get_settings(conn).get("base_url"):
                D.set_setting(conn, "base_url", f"http://{ip}:8000")
        print("\n" + "=" * 60, flush=True)
        print("  DAVOMAT TIZIMI ishga tushdi", flush=True)
        print(f"  Proyektor ekrani : http://{ip}:8000/projector", flush=True)
        print(f"  Admin panel      : http://{ip}:8000/admin", flush=True)
        print(f"  Xodimlar QR orqali shu manzilga tushadi: http://{ip}:8000", flush=True)
        print("=" * 60 + "\n", flush=True)
    except Exception as exc:  # noqa: BLE001 — server baribir ko'tarilsin
        print(f"\n  OGOHLANTIRISH: baza tayyorlanmadi — {exc}\n", flush=True)
    yield


app = FastAPI(title="Davomat", lifespan=lifespan)
# Lokal ishlaganda statik fayllarni shu mount tarqatadi. Vercel'da esa
# `public/` papkasini CDN o'zi xizmat qiladi va uni funksiya paketiga
# qo'shmasligi mumkin — shuning uchun mavjudligini tekshiramiz.
_static_dir = BASE_DIR / "public" / "static"
if _static_dir.is_dir():
    app.mount("/static", StaticFiles(directory=_static_dir), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "app" / "templates"))

ADMIN_COOKIE = "davomat_admin"
DEVICE_COOKIE = "davomat_device"


# --------------------------------------------------------------------------
# Yordamchi funksiyalar
# --------------------------------------------------------------------------

def lan_ip() -> str:
    """Kompyuterning lokal tarmoqdagi IP manzilini aniqlaydi."""
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        s.close()


def serializer(secret: str) -> URLSafeSerializer:
    return URLSafeSerializer(secret, salt="admin-session")


def is_admin(request: Request, settings: dict) -> bool:
    raw = request.cookies.get(ADMIN_COOKIE)
    if not raw:
        return False
    try:
        data = serializer(settings["qr_secret"]).loads(raw)
    except BadSignature:
        return False
    return data.get("pw") == settings.get("admin_password")


def device_id_of(request: Request) -> str:
    return request.cookies.get(DEVICE_COOKIE) or secrets.token_hex(8)


def client_ip(request: Request) -> str:
    return request.client.host if request.client else ""


def hhmm(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso).strftime("%H:%M")
    except ValueError:
        return ""


def parse_hhmm(value: str, fallback=(9, 0)) -> dtime:
    try:
        h, m = value.split(":")
        return dtime(int(h), int(m))
    except (ValueError, AttributeError):
        return dtime(*fallback)


def first_name(full_name: str) -> str:
    """O'zbekcha F.I.Sh. dan ismni ajratadi: "Aliyev Vali ..." -> "Vali"."""
    parts = full_name.split()
    return parts[1] if len(parts) >= 2 else (parts[0] if parts else "Xodim")


def uz_date(d: date) -> str:
    return f"{d.day}-{MONTHS_UZ[d.month - 1]} {d.year}, {WEEKDAYS_UZ[d.weekday()]}"


def worked_minutes(row) -> int:
    if not row or not row["check_in"] or not row["check_out"]:
        return 0
    try:
        a = datetime.fromisoformat(row["check_in"])
        b = datetime.fromisoformat(row["check_out"])
    except ValueError:
        return 0
    return max(0, int((b - a).total_seconds() // 60))


def fmt_minutes(total: int) -> str:
    if total <= 0:
        return "—"
    return f"{total // 60} soat {total % 60} daq"


templates.env.filters["hhmm"] = hhmm
templates.env.filters["fmt_minutes"] = fmt_minutes
templates.env.globals["STATUS_LABELS"] = STATUS_LABELS
templates.env.globals["uz_date"] = uz_date


def auto_base_url(request: Request) -> str:
    """Manzilni so'rovdan aniqlaydi: bulutda domen, lokalda LAN IP."""
    if IS_CLOUD:
        # Vercel orqasida turganda haqiqiy domen sarlavhalarda keladi
        host = request.headers.get("x-forwarded-host") or request.headers.get("host")
        scheme = request.headers.get("x-forwarded-proto", "https")
        if host:
            return f"{scheme}://{host}"
    port = request.url.port or 8000
    return f"http://{lan_ip()}:{port}"


def base_url(conn, request: Request) -> str:
    """QR koddagi manzil: qo'lda kiritilgani, bo'lmasa avtomatik aniqlangani."""
    settings = D.get_settings(conn)
    custom = (settings.get("base_url") or "").strip().rstrip("/")
    return custom or auto_base_url(request)


def distance_m(lat1: float, lng1: float, lat2: float, lng2: float) -> float:
    """Ikki koordinata orasidagi masofa, metrda (haversine)."""
    r = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lng2 - lng1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def geo_config(settings: dict):
    """Joylashuv tekshiruvi yoqilganmi va ofis koordinatasi bormi?"""
    if settings.get("geo_required") != "1":
        return None
    try:
        lat = float(settings.get("office_lat") or "")
        lng = float(settings.get("office_lng") or "")
    except ValueError:
        return None
    try:
        radius = float(settings.get("geo_radius_m") or "200")
    except ValueError:
        radius = 200.0
    return lat, lng, radius


def error_page(title: str, body_html: str, status: int) -> HTMLResponse:
    """Oddiy, tushunarli xato sahifasi (shablonlarga bog'liq emas)."""
    return HTMLResponse(
        f"<!doctype html><html lang=uz><meta charset=utf-8><title>{title}</title>"
        "<body style=\"font-family:system-ui,-apple-system,Segoe UI,sans-serif;"
        "background:#0b0d12;color:#e8ecf4;display:grid;place-items:center;"
        "min-height:100vh;margin:0;padding:24px\">"
        "<div style='max-width:540px;text-align:center'>"
        f"<h1 style='font-size:24px;margin:0 0 14px'>{title}</h1>{body_html}"
        "</div></body></html>",
        status_code=status,
    )


CODE = ("background:#1e2432;padding:2px 7px;border-radius:5px;"
        "font-family:ui-monospace,monospace")
DIM = "color:#8b95aa;line-height:1.7;margin:0 0 12px"


@app.exception_handler(Exception)
async def unhandled_error(request: Request, exc: Exception):
    """Kutilmagan xatoni loglaymiz va sababini sahifada ko'rsatamiz."""
    import traceback
    traceback.print_exc()
    hint = ""
    if not D.IS_PG:
        hint = (f"<p style='{DIM}'>Ehtimoliy sabab: "
                f"<code style='{CODE}'>DATABASE_URL</code> o'zgaruvchisi "
                "qo'shilmagan.</p>")
    return error_page(
        "Xatolik yuz berdi",
        f"<p style='{DIM}'>{type(exc).__name__}: {str(exc)[:200]}</p>{hint}"
        f"<p style='{DIM}'>Batafsil ma'lumot Vercel'dagi "
        "<b>Logs</b> bo'limida.</p>",
        500,
    )


@app.middleware("http")
async def require_database(request: Request, call_next):
    """Bulutda baza ulanmagan bo'lsa — tushunarli xabar, stack trace emas.

    Vercel'da disk faqat o'qish uchun, shuning uchun SQLite zaxira varianti
    u yerda ishlay olmaydi: `DATABASE_URL` majburiy.
    """
    if os.environ.get("VERCEL") and not D.IS_PG:
        found = D.db_env_names()
        if found:
            # O'zgaruvchi bor, lekin yaroqli postgres manzili emas
            diag = ("<p style='%s'>Topilgan o'zgaruvchilar: %s</p>"
                    "<p style='%s'>Ulardan hech biri "
                    "<code style='%s'>postgresql://</code> bilan boshlanmaydi — "
                    "qiymat noto'g'ri nusxalangan bo'lishi mumkin.</p>"
                    % (DIM,
                       ", ".join(f"<code style='{CODE}'>{n}</code>" for n in found),
                       DIM, CODE))
        else:
            diag = (f"<p style='{DIM}'>Hozircha bironta ham baza o'zgaruvchisi "
                    "ko'rinmayapti. Tekshiring: o'zgaruvchi <b>Production</b> "
                    "muhiti uchun qo'shilganmi va undan <b>keyin</b> Redeploy "
                    "qilinganmi?</p>")
        return error_page(
            "Ma'lumotlar bazasi ulanmagan",
            f"<p style='{DIM}'>Vercel loyihasida <b>Storage &rarr; Create Database "
            "&rarr; Neon</b> orqali baza ulang, yoki <b>Settings &rarr; Environment "
            f"Variables</b> bo'limiga <code style='{CODE}'>DATABASE_URL</code> "
            "qo'shing (Neon'ning <b>pooled</b> manzili).</p>"
            + diag +
            f"<p style='{DIM}'>Batafsil: repozitoriydagi "
            f"<code style='{CODE}'>VERCEL.md</code></p>",
            503,
        )
    return await call_next(request)


# --------------------------------------------------------------------------
# Proyektor ekrani
# --------------------------------------------------------------------------

@app.get("/", response_class=HTMLResponse)
def root():
    return RedirectResponse("/projector", status_code=302)


@app.get("/projector", response_class=HTMLResponse)
def projector(request: Request):
    with D.db() as conn:
        settings = D.get_settings(conn)
    return templates.TemplateResponse(
        request, "projector.html",
        {
            "request": request,
            "settings": settings,
            "rotate": QR_ROTATE_SECONDS,
            "today_label": uz_date(D.now().date()),
        },
    )


def qr_data_uri(text: str) -> str:
    qr = qrcode.QRCode(version=None, box_size=10, border=2,
                       error_correction=qrcode.constants.ERROR_CORRECT_M)
    qr.add_data(text)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


@app.get("/api/state")
def api_state(request: Request):
    """Proyektor ekrani shu endpointdan yangilanib turadi.

    Bir so'rovda bir nechta kelgusi QR ham qaytariladi — ekran ularni o'zi
    almashtirib turadi va serverga kamdan-kam murojaat qiladi.
    """
    with D.db() as conn:
        settings = D.get_settings(conn)
        root = base_url(conn, request)
        win = current_window()
        slots = [
            {
                "window": win + i,
                "token": (tok := make_token(settings["qr_secret"], win + i)),
                "qr": qr_data_uri(f"{root}/c/{tok}"),
            }
            for i in range(QR_SLOTS_AHEAD)
        ]
        token = slots[0]["token"]
        url = f"{root}/c/{token}"
        today = D.today_str()

        rows = conn.execute(
            "SELECT a.*, e.full_name FROM attendance a "
            "JOIN employees e ON e.id = a.employee_id "
            "WHERE a.work_date = ? AND a.check_in IS NOT NULL "
            "ORDER BY a.check_in DESC",
            (today,),
        ).fetchall()
        total = conn.execute(
            "SELECT COUNT(*) c FROM employees WHERE active = 1"
        ).fetchone()["c"]

    recent = [
        {
            "name": r["full_name"],
            "time": hhmm(r["check_in"]),
            "out": hhmm(r["check_out"]),
            "late": r["late_minutes"],
        }
        for r in rows
    ]
    return JSONResponse({
        "token": token,
        "url": url,
        "qr": slots[0]["qr"],
        "slots": slots,
        "window": win,
        "seconds_left": seconds_left(),
        "rotate": QR_ROTATE_SECONDS,
        "clock": D.now().strftime("%H:%M:%S"),
        "present": len(recent),
        "total": total,
        "recent": recent,
    })


# --------------------------------------------------------------------------
# Telefon sahifasi — QR skanerlangandan keyin
# --------------------------------------------------------------------------

@app.get("/c/{token}", response_class=HTMLResponse)
def checkin_page(request: Request, token: str):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if not verify_token(settings["qr_secret"], token):
            return templates.TemplateResponse(
                request, "message.html",
                {"request": request, "settings": settings, "ok": False,
                 "title": "QR kod eskirgan",
                 "text": "Ekrandagi yangi QR kodni qaytadan skanerlang."},
                status_code=410,
            )

        today = D.today_str()
        employees = D.active_employees(conn)
        records = {
            r["employee_id"]: r
            for r in conn.execute(
                "SELECT * FROM attendance WHERE work_date = ?", (today,)
            ).fetchall()
        }
        absences = {
            a["employee_id"]: a
            for a in conn.execute(
                "SELECT * FROM absences WHERE date_from <= ? AND date_to >= ?",
                (today, today),
            ).fetchall()
        }

    items = []
    for e in employees:
        rec = records.get(e["id"])
        if rec and rec["check_out"]:
            state, action = "ketgan", None
        elif rec and rec["check_in"]:
            state, action = "kelgan", "ketish"
        else:
            state, action = "yoq", "kelish"
        items.append({
            "id": e["id"],
            "name": e["full_name"],
            "position": e["position"],
            "state": state,
            "action": action,
            "check_in": hhmm(rec["check_in"]) if rec else "",
            "check_out": hhmm(rec["check_out"]) if rec else "",
            "absence": STATUS_LABELS.get(absences[e["id"]]["kind"]) if e["id"] in absences else "",
        })

    resp = templates.TemplateResponse(
        request, "checkin.html",
        {"request": request, "settings": settings, "token": token,
         "items": items, "today_label": uz_date(D.now().date()),
         "rotate": QR_ROTATE_SECONDS,
         "geo_required": geo_config(settings) is not None},
    )
    if not request.cookies.get(DEVICE_COOKIE):
        resp.set_cookie(DEVICE_COOKIE, secrets.token_hex(8),
                        max_age=60 * 60 * 24 * 365, samesite="lax")
    return resp


@app.post("/c/{token}/mark", response_class=HTMLResponse)
def mark(request: Request, token: str,
         employee_id: int = Form(...), action: str = Form(...),
         lat: str = Form(""), lng: str = Form(""), acc: str = Form("")):
    device = device_id_of(request)
    ip = client_ip(request)
    ua = request.headers.get("user-agent", "")

    with D.db() as conn:
        settings = D.get_settings(conn)

        def msg(ok: bool, title: str, text: str, status=200):
            resp = templates.TemplateResponse(
                request, "message.html",
                {"request": request, "settings": settings, "ok": ok,
                 "title": title, "text": text},
                status_code=status,
            )
            resp.set_cookie(DEVICE_COOKIE, device,
                            max_age=60 * 60 * 24 * 365, samesite="lax")
            return resp

        if not verify_token(settings["qr_secret"], token):
            return msg(False, "QR kod eskirgan",
                       "Ekrandagi yangi QR kodni qaytadan skanerlang.", 410)

        emp = conn.execute(
            "SELECT * FROM employees WHERE id = ? AND active = 1", (employee_id,)
        ).fetchone()
        if not emp:
            return msg(False, "Xodim topilmadi",
                       "Ro'yxat yangilangan bo'lishi mumkin. Qaytadan urinib ko'ring.", 404)

        # Joylashuv tekshiruvi (yoqilgan bo'lsa)
        geo = geo_config(settings)
        if geo:
            office_lat, office_lng, radius = geo
            try:
                here = distance_m(office_lat, office_lng, float(lat), float(lng))
            except ValueError:
                return msg(False, "Joylashuv aniqlanmadi",
                           "Brauzerga joylashuvga ruxsat bering va QR kodni "
                           "qaytadan skanerlang.")
            # Telefon aniqligini hisobga olamiz: xato chegarasi radiusga qo'shiladi
            try:
                slack = min(float(acc or 0), 150.0)
            except ValueError:
                slack = 0.0
            if here > radius + slack:
                D.log(conn, "geo_rad_etildi", employee_id, device, ip, ua,
                      f"{int(here)}m")
                return msg(False, "Siz ofisda emassiz",
                           f"Ofisdan taxminan {int(here)} metr uzoqdasiz. "
                           f"Belgilash faqat ofis hududida ({int(radius)} m) mumkin.")

        now = D.now()
        today = now.date().isoformat()
        rec = conn.execute(
            "SELECT * FROM attendance WHERE employee_id = ? AND work_date = ?",
            (employee_id, today),
        ).fetchone()

        # Bitta telefondan haddan ziyod ko'p odam belgilanayotgan bo'lsa — bayroqcha
        limit = int(settings.get("device_limit_per_hour", "3"))
        hour_ago = (now - timedelta(hours=1)).isoformat(timespec="seconds")
        distinct = conn.execute(
            "SELECT COUNT(DISTINCT employee_id) c FROM audit "
            "WHERE device_id = ? AND ts >= ? AND action IN ('kelish', 'ketish')",
            (device, hour_ago),
        ).fetchone()["c"]
        flagged = 1 if distinct >= limit else 0

        if action == "kelish":
            if rec and rec["check_in"]:
                return msg(False, "Siz allaqachon belgilangansiz",
                           f"Kelgan vaqtingiz: {hhmm(rec['check_in'])}")

            start = parse_hhmm(settings.get("work_start", "09:00"))
            grace = int(settings.get("late_grace_minutes", "5"))
            deadline = datetime.combine(now.date(), start, tzinfo=TZ) + timedelta(minutes=grace)
            late = 0
            status = "keldi"
            if now > deadline:
                late = int((now - datetime.combine(now.date(), start, tzinfo=TZ)).total_seconds() // 60)
                status = "kechikdi"

            conn.execute(
                "INSERT INTO attendance (employee_id, work_date, check_in, late_minutes,"
                " status, device_id, ip, flagged) VALUES (?, ?, ?, ?, ?, ?, ?, ?)"
                " ON CONFLICT(employee_id, work_date) DO UPDATE SET"
                " check_in = excluded.check_in, late_minutes = excluded.late_minutes,"
                " status = excluded.status, device_id = excluded.device_id,"
                " ip = excluded.ip, flagged = excluded.flagged",
                (employee_id, today, now.isoformat(timespec="seconds"), late,
                 status, device, ip, flagged),
            )
            D.log(conn, "kelish", employee_id, device, ip, ua,
                  f"late={late}")

            if late:
                return msg(True, f"Xush kelibsiz, {first_name(emp['full_name'])}!",
                           f"Kelgan vaqt: {now.strftime('%H:%M')} — {late} daqiqa kechikdingiz.")
            return msg(True, f"Xush kelibsiz, {first_name(emp['full_name'])}!",
                       f"Kelgan vaqt: {now.strftime('%H:%M')}. Yaxshi ish kuni tilaymiz!")

        if action == "ketish":
            if not rec or not rec["check_in"]:
                return msg(False, "Avval kelishni belgilang",
                           "Bugun sizning kelganingiz qayd etilmagan.")
            if rec["check_out"]:
                return msg(False, "Siz allaqachon ketgan deb belgilangansiz",
                           f"Ketgan vaqtingiz: {hhmm(rec['check_out'])}")

            min_shift = int(settings.get("min_shift_minutes", "60"))
            elapsed = int((now - datetime.fromisoformat(rec["check_in"])).total_seconds() // 60)
            if elapsed < min_shift:
                return msg(False, "Juda erta",
                           f"Kelganingizga {elapsed} daqiqa bo'ldi. "
                           f"Ketishni belgilash uchun kamida {min_shift} daqiqa kerak.")

            conn.execute(
                "UPDATE attendance SET check_out = ? WHERE id = ?",
                (now.isoformat(timespec="seconds"), rec["id"]),
            )
            D.log(conn, "ketish", employee_id, device, ip, ua, f"worked={elapsed}")
            return msg(True, "Yaxshi boring!",
                       f"Ketgan vaqt: {now.strftime('%H:%M')}. "
                       f"Bugun ishlagan vaqtingiz: {fmt_minutes(elapsed)}.")

        return msg(False, "Noma'lum amal", "Qaytadan urinib ko'ring.", 400)


# --------------------------------------------------------------------------
# Admin — kirish
# --------------------------------------------------------------------------

@app.get("/admin/login", response_class=HTMLResponse)
def admin_login_page(request: Request, xato: int = 0):
    with D.db() as conn:
        settings = D.get_settings(conn)
    return templates.TemplateResponse(
                request, "login.html", {"request": request, "settings": settings, "xato": xato}
    )


@app.post("/admin/login")
def admin_login(request: Request, password: str = Form(...)):
    with D.db() as conn:
        settings = D.get_settings(conn)
    if password != settings.get("admin_password"):
        return RedirectResponse("/admin/login?xato=1", status_code=303)
    resp = RedirectResponse("/admin", status_code=303)
    token = serializer(settings["qr_secret"]).dumps({"pw": password})
    resp.set_cookie(ADMIN_COOKIE, token, max_age=60 * 60 * 12,
                    httponly=True, samesite="lax")
    return resp


@app.get("/admin/logout")
def admin_logout():
    resp = RedirectResponse("/admin/login", status_code=303)
    resp.delete_cookie(ADMIN_COOKIE)
    return resp


def guard(request: Request, settings: dict):
    if not is_admin(request, settings):
        return RedirectResponse("/admin/login", status_code=303)
    return None


# --------------------------------------------------------------------------
# Admin — bugungi holat
# --------------------------------------------------------------------------

@app.get("/admin", response_class=HTMLResponse)
def admin_home(request: Request, kun: str = ""):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r

        day = D.parse_date(kun) if kun else D.now().date()
        day_s = day.isoformat()

        employees = D.active_employees(conn)
        records = {
            r["employee_id"]: r
            for r in conn.execute(
                "SELECT * FROM attendance WHERE work_date = ?", (day_s,)
            ).fetchall()
        }
        absences = {
            a["employee_id"]: a
            for a in conn.execute(
                "SELECT * FROM absences WHERE date_from <= ? AND date_to >= ?",
                (day_s, day_s),
            ).fetchall()
        }

    rows, stats = [], {"keldi": 0, "kechikdi": 0, "kelmadi": 0, "sababli": 0}
    for e in employees:
        rec = records.get(e["id"])
        ab = absences.get(e["id"])
        if rec and rec["check_in"]:
            status = rec["status"]
        elif ab:
            status = ab["kind"]
        else:
            status = "kelmadi"

        if status in ("keldi", "kechikdi", "kelmadi"):
            stats[status] += 1
        else:
            stats["sababli"] += 1

        rows.append({
            "id": e["id"],
            "name": e["full_name"],
            "position": e["position"],
            "department": e["department"],
            "status": status,
            "check_in": hhmm(rec["check_in"]) if rec else "",
            "check_out": hhmm(rec["check_out"]) if rec else "",
            "late": rec["late_minutes"] if rec else 0,
            "worked": worked_minutes(rec),
            "flagged": rec["flagged"] if rec else 0,
            "manual": rec["manual"] if rec else 0,
            "note": (rec["note"] if rec else "") or (ab["note"] if ab else ""),
        })

    return templates.TemplateResponse(
        request, "admin_today.html",
        {"request": request, "settings": settings, "rows": rows, "stats": stats,
         "day": day_s, "day_label": uz_date(day),
         "prev_day": (day - timedelta(days=1)).isoformat(),
         "next_day": (day + timedelta(days=1)).isoformat(),
         "today": D.today_str(), "active": "today"},
    )


@app.post("/admin/attendance/manual")
def manual_edit(request: Request, employee_id: int = Form(...), day: str = Form(...),
                check_in: str = Form(""), check_out: str = Form(""),
                note: str = Form("")):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r

        d = D.parse_date(day)

        def to_iso(value: str):
            if not value:
                return None
            t = parse_hhmm(value, (0, 0))
            return datetime.combine(d, t, tzinfo=TZ).isoformat(timespec="seconds")

        ci, co = to_iso(check_in), to_iso(check_out)
        if not ci:
            conn.execute(
                "DELETE FROM attendance WHERE employee_id = ? AND work_date = ?",
                (employee_id, d.isoformat()),
            )
            D.log(conn, "admin_ochirdi", employee_id, detail=d.isoformat())
        else:
            start = parse_hhmm(settings.get("work_start", "09:00"))
            grace = int(settings.get("late_grace_minutes", "5"))
            in_dt = datetime.fromisoformat(ci)
            deadline = datetime.combine(d, start, tzinfo=TZ) + timedelta(minutes=grace)
            late = 0
            status = "keldi"
            if in_dt > deadline:
                late = int((in_dt - datetime.combine(d, start, tzinfo=TZ)).total_seconds() // 60)
                status = "kechikdi"
            conn.execute(
                "INSERT INTO attendance (employee_id, work_date, check_in, check_out,"
                " late_minutes, status, note, manual) VALUES (?, ?, ?, ?, ?, ?, ?, 1)"
                " ON CONFLICT(employee_id, work_date) DO UPDATE SET"
                " check_in = excluded.check_in, check_out = excluded.check_out,"
                " late_minutes = excluded.late_minutes, status = excluded.status,"
                " note = excluded.note, manual = 1, flagged = 0",
                (employee_id, d.isoformat(), ci, co, late, status, note),
            )
            D.log(conn, "admin_tahrirladi", employee_id, detail=d.isoformat())

    return RedirectResponse(f"/admin?kun={day}", status_code=303)


# --------------------------------------------------------------------------
# Admin — xodimlar
# --------------------------------------------------------------------------

@app.get("/admin/xodimlar", response_class=HTMLResponse)
def employees_page(request: Request):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        rows = D.all_employees(conn)
    return templates.TemplateResponse(
        request, "admin_employees.html",
        {"request": request, "settings": settings, "rows": rows, "active": "employees"},
    )


@app.post("/admin/xodimlar/qoshish")
def employee_add(request: Request, full_name: str = Form(...),
                 position: str = Form(""), department: str = Form(""),
                 phone: str = Form("")):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        name = full_name.strip()
        if name:
            conn.execute(
                "INSERT INTO employees (full_name, position, department, phone, created_at)"
                " VALUES (?, ?, ?, ?, ?)",
                (name, position.strip(), department.strip(), phone.strip(),
                 D.now().isoformat(timespec="seconds")),
            )
    return RedirectResponse("/admin/xodimlar", status_code=303)


@app.post("/admin/xodimlar/{emp_id}/tahrir")
def employee_edit(request: Request, emp_id: int, full_name: str = Form(...),
                  position: str = Form(""), department: str = Form(""),
                  phone: str = Form("")):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        conn.execute(
            "UPDATE employees SET full_name = ?, position = ?, department = ?, phone = ?"
            " WHERE id = ?",
            (full_name.strip(), position.strip(), department.strip(), phone.strip(), emp_id),
        )
    return RedirectResponse("/admin/xodimlar", status_code=303)


@app.post("/admin/xodimlar/{emp_id}/holat")
def employee_toggle(request: Request, emp_id: int):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        conn.execute("UPDATE employees SET active = 1 - active WHERE id = ?", (emp_id,))
    return RedirectResponse("/admin/xodimlar", status_code=303)


@app.post("/admin/xodimlar/{emp_id}/ochirish")
def employee_delete(request: Request, emp_id: int):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        conn.execute("DELETE FROM employees WHERE id = ?", (emp_id,))
    return RedirectResponse("/admin/xodimlar", status_code=303)


# --------------------------------------------------------------------------
# Admin — sababli yo'qliklar
# --------------------------------------------------------------------------

@app.get("/admin/yoqliklar", response_class=HTMLResponse)
def absences_page(request: Request):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        raw = conn.execute(
            "SELECT ab.*, e.full_name FROM absences ab"
            " JOIN employees e ON e.id = ab.employee_id"
            " ORDER BY ab.date_from DESC"
        ).fetchall()
        employees = D.active_employees(conn)

    rows = [
        dict(r, days=(D.parse_date(r["date_to"]) - D.parse_date(r["date_from"])).days + 1)
        for r in raw
    ]
    return templates.TemplateResponse(
        request, "admin_absences.html",
        {"request": request, "settings": settings, "rows": rows,
         "employees": employees, "kinds": ABSENCE_TYPES,
         "today": D.today_str(), "active": "absences"},
    )


@app.post("/admin/yoqliklar/qoshish")
def absence_add(request: Request, employee_id: int = Form(...),
                date_from: str = Form(...), date_to: str = Form(...),
                kind: str = Form(...), note: str = Form("")):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        a, b = D.parse_date(date_from), D.parse_date(date_to)
        if b < a:
            a, b = b, a
        if kind in ABSENCE_TYPES:
            conn.execute(
                "INSERT INTO absences (employee_id, date_from, date_to, kind, note, created_at)"
                " VALUES (?, ?, ?, ?, ?, ?)",
                (employee_id, a.isoformat(), b.isoformat(), kind, note.strip(),
                 D.now().isoformat(timespec="seconds")),
            )
    return RedirectResponse("/admin/yoqliklar", status_code=303)


@app.post("/admin/yoqliklar/{ab_id}/ochirish")
def absence_delete(request: Request, ab_id: int):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        conn.execute("DELETE FROM absences WHERE id = ?", (ab_id,))
    return RedirectResponse("/admin/yoqliklar", status_code=303)


# --------------------------------------------------------------------------
# Admin — hisobot
# --------------------------------------------------------------------------

def build_report(conn, start: date, end: date):
    """Davr uchun har bir xodim kesimida yig'ma hisobot."""
    employees = D.active_employees(conn)
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    workdays = [d for d in days if d.weekday() < 5]

    att = {}
    for r in conn.execute(
        "SELECT * FROM attendance WHERE work_date BETWEEN ? AND ?",
        (start.isoformat(), end.isoformat()),
    ).fetchall():
        att[(r["employee_id"], r["work_date"])] = r

    absences = conn.execute(
        "SELECT * FROM absences WHERE date_from <= ? AND date_to >= ?",
        (end.isoformat(), start.isoformat()),
    ).fetchall()

    def absent_kind(emp_id: int, d: str):
        for a in absences:
            if a["employee_id"] == emp_id and a["date_from"] <= d <= a["date_to"]:
                return a["kind"]
        return None

    report = []
    for e in employees:
        came = late = missed = excused = 0
        late_total = worked_total = 0
        cells = []
        for d in days:
            ds = d.isoformat()
            rec = att.get((e["id"], ds))
            kind = absent_kind(e["id"], ds)
            weekend = d.weekday() >= 5
            if rec and rec["check_in"]:
                status = rec["status"]
                if status == "kechikdi":
                    late += 1
                    late_total += rec["late_minutes"]
                else:
                    came += 1
                worked_total += worked_minutes(rec)
            elif kind:
                status = kind
                excused += 1
            elif weekend:
                status = "dam"
            else:
                status = "kelmadi"
                missed += 1
            cells.append({
                "date": ds, "day": d.day, "status": status, "weekend": weekend,
                "check_in": hhmm(rec["check_in"]) if rec else "",
                "check_out": hhmm(rec["check_out"]) if rec else "",
                "late": rec["late_minutes"] if rec else 0,
            })

        report.append({
            "id": e["id"], "name": e["full_name"], "position": e["position"],
            "department": e["department"], "cells": cells,
            "came": came, "late": late, "missed": missed, "excused": excused,
            "late_total": late_total, "worked_total": worked_total,
            "present_days": came + late,
            "workdays": len(workdays),
        })
    return report, days


@app.get("/admin/hisobot", response_class=HTMLResponse)
def report_page(request: Request, boshi: str = "", oxiri: str = ""):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r

        today = D.now().date()
        start = D.parse_date(boshi, today.replace(day=1)) if boshi else today.replace(day=1)
        end = D.parse_date(oxiri, today) if oxiri else today
        if end < start:
            start, end = end, start
        if (end - start).days > 92:
            end = start + timedelta(days=92)

        report, days = build_report(conn, start, end)

    return templates.TemplateResponse(
        request, "admin_report.html",
        {"request": request, "settings": settings, "report": report, "days": days,
         "start": start.isoformat(), "end": end.isoformat(),
         "period_label": f"{uz_date(start)} — {uz_date(end)}",
         "active": "report"},
    )


@app.get("/admin/hisobot/excel")
def report_excel(request: Request, boshi: str = "", oxiri: str = ""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
        today = D.now().date()
        start = D.parse_date(boshi, today.replace(day=1)) if boshi else today.replace(day=1)
        end = D.parse_date(oxiri, today) if oxiri else today
        if end < start:
            start, end = end, start
        report, days = build_report(conn, start, end)

    wb = Workbook()

    # 1-varaq: yig'ma
    ws = wb.active
    ws.title = "Yig'ma"
    headers = ["F.I.Sh.", "Lavozim", "Bo'lim", "Keldi", "Kechikdi",
               "Kelmadi", "Sababli", "Jami kechikish (daq)", "Ishlangan vaqt (soat)"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in report:
        ws.append([r["name"], r["position"], r["department"], r["came"], r["late"],
                   r["missed"], r["excused"], r["late_total"],
                   round(r["worked_total"] / 60, 1)])
    for col, width in zip("ABCDEFGHI", [28, 20, 18, 9, 11, 11, 10, 20, 22]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    # 2-varaq: kunlik jadval
    ws2 = wb.create_sheet("Kunlik")
    head2 = ["F.I.Sh."] + [f"{d.day:02d}.{d.month:02d}" for d in days]
    ws2.append(head2)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(horizontal="center")
    colors = {"keldi": "D1FAE5", "kechikdi": "FEF3C7", "kelmadi": "FEE2E2",
              "dam": "F3F4F6"}
    for r in report:
        row = [r["name"]]
        for cell in r["cells"]:
            if cell["status"] in ("keldi", "kechikdi"):
                txt = cell["check_in"] + (f"–{cell['check_out']}" if cell["check_out"] else "")
            elif cell["status"] == "dam":
                txt = "—"
            else:
                txt = STATUS_LABELS.get(cell["status"], cell["status"])
            row.append(txt)
        ws2.append(row)
        for i, cell in enumerate(r["cells"], start=2):
            fill = colors.get(cell["status"], "E0E7FF")
            ws2.cell(row=ws2.max_row, column=i).fill = PatternFill("solid", fgColor=fill)
            ws2.cell(row=ws2.max_row, column=i).alignment = Alignment(horizontal="center")
    ws2.column_dimensions["A"].width = 28
    for i in range(2, len(days) + 2):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = 13
    ws2.freeze_panes = "B2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"davomat_{start.isoformat()}_{end.isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# --------------------------------------------------------------------------
# Admin — sozlamalar
# --------------------------------------------------------------------------

@app.get("/admin/sozlamalar", response_class=HTMLResponse)
def settings_page(request: Request, saqlandi: int = 0):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r
    return templates.TemplateResponse(
        request, "admin_settings.html",
        {"request": request, "settings": settings, "saqlandi": saqlandi,
         "lan_ip": "" if IS_CLOUD else lan_ip(), "is_cloud": IS_CLOUD,
         "auto_url": auto_base_url(request), "active": "settings"},
    )


@app.post("/admin/sozlamalar")
def settings_save(request: Request, company_name: str = Form(""),
                  work_start: str = Form("09:00"), work_end: str = Form("18:00"),
                  late_grace_minutes: str = Form("5"),
                  min_shift_minutes: str = Form("60"),
                  device_limit_per_hour: str = Form("3"),
                  base_url_value: str = Form(""),
                  geo_required: str = Form(""),
                  office_lat: str = Form(""), office_lng: str = Form(""),
                  geo_radius_m: str = Form("200"),
                  admin_password: str = Form("")):
    with D.db() as conn:
        settings = D.get_settings(conn)
        if (r := guard(request, settings)):
            return r

        def num(value: str, fallback: str) -> str:
            return value if value.isdigit() else fallback

        D.set_setting(conn, "company_name", company_name.strip() or "Kompaniya")
        D.set_setting(conn, "work_start", work_start or "09:00")
        D.set_setting(conn, "work_end", work_end or "18:00")
        D.set_setting(conn, "late_grace_minutes", num(late_grace_minutes, "5"))
        D.set_setting(conn, "min_shift_minutes", num(min_shift_minutes, "60"))
        D.set_setting(conn, "device_limit_per_hour", num(device_limit_per_hour, "3"))
        D.set_setting(conn, "base_url", base_url_value.strip().rstrip("/"))

        def coord(value: str) -> str:
            value = value.strip().replace(",", ".")
            try:
                float(value)
            except ValueError:
                return ""
            return value

        D.set_setting(conn, "office_lat", coord(office_lat))
        D.set_setting(conn, "office_lng", coord(office_lng))
        D.set_setting(conn, "geo_radius_m", num(geo_radius_m, "200"))
        D.set_setting(conn, "geo_required", "1" if geo_required else "0")

        if admin_password.strip():
            D.set_setting(conn, "admin_password", admin_password.strip())

    resp = RedirectResponse("/admin/sozlamalar?saqlandi=1", status_code=303)
    if admin_password.strip():
        resp.delete_cookie(ADMIN_COOKIE)
    return resp
